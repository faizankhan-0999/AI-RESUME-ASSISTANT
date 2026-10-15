"""AI Resume ATS Checker

Upload a resume (PDF, DOCX or TXT) and get an ATS score plus concrete
improvements, powered by a Gemini Flash model and a Streamlit UI.
"""

import io
import os
from typing import List, Optional

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pydantic import BaseModel, Field
from pypdf import PdfReader

DEFAULT_MODEL = "gemini-3.8-flash"
MAX_RESUME_CHARS = 30_000
MAX_JD_CHARS = 10_000
MIN_RESUME_CHARS = 200


# ---------------------------------------------------------------------------
# Structured output schema (what we ask Gemini to return)
# ---------------------------------------------------------------------------
class CategoryScore(BaseModel):
    name: str
    score: int = Field(description="0-100")
    comment: str


class Improvement(BaseModel):
    priority: str = Field(description="High, Medium or Low")
    section: str
    issue: str
    suggestion: str


class RewriteExample(BaseModel):
    original: str
    improved: str


class ATSReport(BaseModel):
    overall_score: int = Field(description="0-100")
    summary: str
    category_scores: List[CategoryScore]
    strengths: List[str]
    missing_keywords: List[str]
    improvements: List[Improvement]
    rewrite_examples: List[RewriteExample]


# ---------------------------------------------------------------------------
# File parsing
# ---------------------------------------------------------------------------
def extract_text(uploaded_file) -> str:
    """Return plain text from an uploaded PDF, DOCX or TXT file."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# ---------------------------------------------------------------------------
# Gemini call
# ---------------------------------------------------------------------------
SYSTEM_INSTRUCTION = (
    "You are an expert technical recruiter and ATS (Applicant Tracking System) "
    "specialist. You evaluate resumes strictly and honestly. Never invent "
    "experience the candidate does not have. Base every statement on the "
    "resume text provided."
)


def build_prompt(resume_text: str, job_description: str) -> str:
    jd_block = (
        f"JOB DESCRIPTION:\n\"\"\"\n{job_description}\n\"\"\"\n\n"
        "Score keyword match against this job description and list the important "
        "keywords/skills from it that are missing from the resume."
        if job_description
        else "No job description was provided. Evaluate the resume for general ATS "
        "friendliness and list commonly expected keywords for the candidate's "
        "apparent target role that are missing."
    )
    return f"""Analyze the resume below for ATS compatibility.

Scoring rules:
- overall_score is an integer 0-100.
- category_scores must contain exactly these 5 categories, each scored 0-100:
  "Keywords & Relevance", "Formatting & Structure", "Impact & Achievements",
  "Clarity & Grammar", "Contact & Completeness".
- strengths: 3-6 short bullet strings.
- missing_keywords: up to 15 concise keywords or skills.
- improvements: 6-10 items ordered by priority (High first). Each has priority
  (High/Medium/Low), the resume section, the issue, and an actionable suggestion.
- rewrite_examples: 2-4 weak bullet points copied from the resume with a stronger
  rewrite using action verbs and measurable results. Do not fabricate numbers;
  use placeholders like [X%] where a metric is needed.
- summary: 2-3 sentences.

{jd_block}

RESUME:
\"\"\"
{resume_text}
\"\"\"
"""


def analyze_resume(
    api_key: str, model: str, resume_text: str, job_description: str
) -> ATSReport:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_INSTRUCTION,
            temperature=0.2,
            response_mime_type="application/json",
            response_schema=ATSReport,
        ),
    )
    text = (response.text or "").strip()
    if not text:
        raise RuntimeError("The model returned an empty response. Please try again.")

    report = ATSReport.model_validate_json(text)
    # Clamp scores to a sane range in case the model drifts.
    report.overall_score = max(0, min(100, report.overall_score))
    for cat in report.category_scores:
        cat.score = max(0, min(100, cat.score))
    return report


# ---------------------------------------------------------------------------
# UI helpers
# ---------------------------------------------------------------------------
def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


def get_api_key() -> Optional[str]:
    """Look for the key in Streamlit secrets, then the environment."""
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass
    return os.environ.get("GEMINI_API_KEY")


def render_report(report: ATSReport) -> None:
    st.divider()
    left, right = st.columns([1, 2])
    with left:
        st.metric("ATS Score", f"{report.overall_score}/100", score_label(report.overall_score))
        st.progress(report.overall_score / 100)
    with right:
        st.subheader("Summary")
        st.write(report.summary)

    st.subheader("Score breakdown")
    cols = st.columns(len(report.category_scores) or 1)
    for col, cat in zip(cols, report.category_scores):
        with col:
            st.metric(cat.name, f"{cat.score}/100")
            st.progress(cat.score / 100)
            st.caption(cat.comment)

    tab_improve, tab_keywords, tab_strengths, tab_rewrite = st.tabs(
        ["Improvements", "Missing keywords", "Strengths", "Bullet rewrites"]
    )

    with tab_improve:
        order = {"high": 0, "medium": 1, "low": 2}
        items = sorted(report.improvements, key=lambda i: order.get(i.priority.lower(), 3))
        icons = {"high": "🔴", "medium": "🟠", "low": "🟢"}
        for item in items:
            icon = icons.get(item.priority.lower(), "⚪")
            with st.expander(f"{icon} {item.priority} · {item.section}: {item.issue}"):
                st.write(item.suggestion)

    with tab_keywords:
        if report.missing_keywords:
            st.write("Consider adding these where they honestly apply to your experience:")
            st.write(" ".join(f"`{k}`" for k in report.missing_keywords))
        else:
            st.success("No major missing keywords found.")

    with tab_strengths:
        for s in report.strengths:
            st.markdown(f"- {s}")

    with tab_rewrite:
        if report.rewrite_examples:
            for ex in report.rewrite_examples:
                st.markdown(f"**Before:** {ex.original}")
                st.markdown(f"**After:** {ex.improved}")
                st.write("")
        else:
            st.info("No rewrite examples were generated.")


# ---------------------------------------------------------------------------
# App
# ---------------------------------------------------------------------------
def main() -> None:
    st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="wide")
    st.title("📄 AI Resume ATS Checker")
    st.caption("Upload your resume to get an ATS score and specific ways to improve it.")

    with st.sidebar:
        st.header("Settings")
        api_key = get_api_key()
        if api_key:
            st.success("Gemini API key loaded.")
        else:
            api_key = st.text_input(
                "Gemini API key",
                type="password",
                help="Get a free key at https://aistudio.google.com/apikey",
            )
        model = st.text_input("Gemini model", value=DEFAULT_MODEL)
        st.caption("Your resume is sent to the Gemini API for analysis and is not stored by this app.")

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional)",
        height=150,
        placeholder="Paste a job description to get a role-specific keyword match...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar.")
            st.stop()

        try:
            with st.spinner("Reading your resume..."):
                resume_text = extract_text(uploaded)
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            st.stop()

        if len(resume_text) < MIN_RESUME_CHARS:
            st.error(
                "Very little text could be extracted. If your resume is a scanned image "
                "or a design-heavy PDF, an ATS would likely struggle to read it too. "
                "Export a text-based PDF or upload a DOCX."
            )
            st.stop()

        try:
            with st.spinner("Analyzing with Gemini..."):
                report = analyze_resume(
                    api_key,
                    model.strip() or DEFAULT_MODEL,
                    resume_text[:MAX_RESUME_CHARS],
                    job_description.strip()[:MAX_JD_CHARS],
                )
            st.session_state["report"] = report
        except Exception as exc:
            st.session_state.pop("report", None)
            st.error(f"Analysis failed: {exc}")
            st.stop()

    report = st.session_state.get("report")
    if report:
        render_report(report)


if __name__ == "__main__":
    main()
